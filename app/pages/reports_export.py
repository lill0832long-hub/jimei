"""报表导出 — Excel/PDF 导出功能"""
import os
from datetime import datetime
from nicegui import ui
from app.components.state import state
from app.components.ui_helpers import show_toast, refresh_main
from app.services import ReportService


def _build_pdf(filepath, title, data, headers, row_fn):
    """通用 PDF 导出（使用 reportlab）"""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_paths = [
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    ]
    font_name = "Helvetica"
    for fp in font_paths:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("CJK", fp))
                font_name = "CJK"
                break
            except Exception:
                pass
    doc = SimpleDocTemplate(filepath, pagesize=landscape(A4),
                            leftMargin=10*mm, rightMargin=10*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontName=font_name, fontSize=16)
    elements = []
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 5*mm))
    table_data = [headers]
    for item in data:
        table_data.append(row_fn(item))
    col_width = (landscape(A4)[0] - 20*mm) / len(headers)
    table = Table(table_data, colWidths=[col_width]*len(headers), repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0,0), (-1,-1), font_name),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1677FF")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#E5E7EB")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    elements.append(table)
    doc.build(elements)


def _export_balance_sheet():
    """导出资产负债表为Excel"""
    try:
        lid = state.selected_ledger_id
        export_dir = os.path.join(os.path.dirname(__file__), "exports")
        os.makedirs(export_dir, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"资产负债表_{state.selected_year}{state.selected_month:02d}_{ts}.xlsx"
        fpath = os.path.join(export_dir, fname)
        report = ReportService.get_balance_sheet(lid, state.selected_year, state.selected_month)
        if not report:
            show_toast("无资产负债表数据", "warning")
            return
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "资产负债表"
        headers = ["科目代码", "科目名称", "期初余额", "期末余额", "变动金额", "变动比例"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        thin = Side(style="thin", color="E5E7EB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        # report 是字典，包含 assets/liabilities/equity 三个列表
        for section_key in ("assets", "liabilities", "equity"):
            for row in report.get(section_key, []):
                data_row = [row.get("code",""), row.get("name",""),
                           row.get("opening_balance",0) or 0, row.get("end",0) or 0,
                           row.get("change",0) or 0, row.get("change_pct","")]
                ws.append(data_row)
            for cell in ws[ws.max_row]:
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
        for col, width in zip("ABCDEF", [12, 20, 14, 14, 14, 12]):
            ws.column_dimensions[col].width = width
        wb.save(fpath)
        show_toast(f"✅ 资产负债表已导出 → {fname}", "success")
        refresh_main()
    except Exception as e:
        show_toast(f"❌ Excel导出失败: {e}", "error")


def _export_balance_sheet_pdf():
    """导出资产负债表为PDF"""
    try:
        lid = state.selected_ledger_id
        export_dir = os.path.join(os.path.dirname(__file__), "exports")
        os.makedirs(export_dir, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"资产负债表_{state.selected_year}{state.selected_month:02d}_{ts}.pdf"
        fpath = os.path.join(export_dir, fname)
        report = ReportService.get_balance_sheet(lid, state.selected_year, state.selected_month)
        if not report:
            show_toast("无资产负债表数据", "warning")
            return
        # report 是字典，包含 assets/liabilities/equity 三个列表
        pdf_data = []
        for section_key in ("assets", "liabilities", "equity"):
            pdf_data.extend(report.get(section_key, []))
        _build_pdf(fpath, "资产负债表", pdf_data,
                   ["科目代码","科目名称","期初余额","期末余额","变动金额","变动比例"],
                   lambda r: [r.get("code",""), r.get("name",""),
                              f'¥{r.get("opening_balance",0) or 0:,.2f}',
                              f'¥{r.get("end",0) or 0:,.2f}',
                              f'¥{r.get("change",0) or 0:,.2f}',
                              r.get("change_pct","")])
        show_toast(f"✅ 资产负债表PDF已导出 → {fname}", "success")
        refresh_main()
    except Exception as e:
        show_toast(f"❌ PDF导出失败: {e}", "error")


def _export_income_statement():
    """导出利润表为Excel"""
    try:
        lid = state.selected_ledger_id
        export_dir = os.path.join(os.path.dirname(__file__), "exports")
        os.makedirs(export_dir, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"利润表_{state.selected_year}{state.selected_month:02d}_{ts}.xlsx"
        fpath = os.path.join(export_dir, fname)
        report = ReportService.get_income_statement(lid, state.selected_year, state.selected_month)
        if not report:
            show_toast("无利润表数据", "warning")
            return
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "利润表"
        headers = ["科目代码", "科目名称", "类别", "本期金额", "上期金额", "变动比例"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        thin = Side(style="thin", color="E5E7EB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        # report 是字典，包含 rows 列表
        for row in report.get("rows", []):
            data_row = [row.get("code",""), row.get("name",""), row.get("category",""),
                       row.get("balance",0) or 0, row.get("prev_balance",0) or 0,
                       row.get("change_pct","")]
            ws.append(data_row)
            for cell in ws[ws.max_row]:
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
        for col, width in zip("ABCDEF", [12, 20, 10, 14, 14, 12]):
            ws.column_dimensions[col].width = width
        wb.save(fpath)
        show_toast(f"✅ 利润表已导出 → {fname}", "success")
        refresh_main()
    except Exception as e:
        show_toast(f"❌ Excel导出失败: {e}", "error")
