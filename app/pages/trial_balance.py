"""科目余额表 — 左右网格布局（无重复年月）"""
from nicegui import ui
from app.components.state import state
from app.components.ui_helpers import format_amount, show_toast, navigate, refresh_main, render_kpi_cards
from app.components.ui_components import SectionHeader, EmptyState
from app.services import LedgerService, ReportService


def render_trial_balance():
    """科目余额表 — 左右网格布局"""
    if not state.selected_ledger_id:
        ledgers = LedgerService.get_all()
        if ledgers:
            state.selected_ledger_id = ledgers[0].get("id") if isinstance(ledgers[0], dict) else ledgers[0]["id"]
    lid = state.selected_ledger_id
    if not lid:
        with ui.card().classes("report-card"):
            with ui.column().classes("report-empty"):
                ui.label("📊").classes("report-empty__icon")
                ui.label("请先创建账套").classes("report-empty__text")
        return

    year, month = state.selected_year, state.selected_month

    # ── 获取数据 ──
    balances = ReportService.get_account_balances(lid, year, month)
    if not balances:
        with ui.card().classes("report-card"):
            EmptyState(icon="grid_on", message="暂无科目余额数据",
                      hint="请先录入记账凭证并过账",
                      action=lambda: navigate("journal"), action_label="去填凭证")
        return

    # ── 分类汇总 ──
    assets, liabilities, equity, revenue, expense = [], [], [], [], []
    for b in balances:
        cat = b.get("category", "")
        if cat == "资产": assets.append(b)
        elif cat == "负债": liabilities.append(b)
        elif cat == "权益": equity.append(b)
        elif cat == "收入": revenue.append(b)
        elif cat == "费用": expense.append(b)

    def _sum_field(items, field):
        return sum(float(b.get(field, 0) if b.get(field) is not None else 0) for b in items)

    total_opening = _sum_field(balances, "opening_balance")
    total_debit = _sum_field(balances, "period_debit")
    total_credit = _sum_field(balances, "period_credit")
    total_closing = _sum_field(balances, "closing_balance")
    diff = abs(total_debit - total_credit)

    # ── KPI 卡片 ──
    kpis = [
        ("期初合计", total_opening, "account_balance", "blue", ""),
        ("本期借方", total_debit, "add_circle", "red", ""),
        ("本期贷方", total_credit, "remove_circle", "green", ""),
        ("期末合计", total_closing, "account_balance_wallet", "purple", ""),
    ]
    if diff >= 0.01:
        kpis.append(("差额", diff, "warning", "orange", "不平衡"))
    render_kpi_cards(kpis)

    # ── 左右网格布局 ──
    # 过滤零余额科目
    _nz = lambda items: [b for b in items if float(b.get('closing_balance', 0) or 0) != 0 or float(b.get('opening_balance', 0) or 0) != 0 or float(b.get('period_debit', 0) or 0) != 0 or float(b.get('period_credit', 0) or 0) != 0]
    _CATEGORY_CONFIG = [
        ("资产类", _nz(assets), "#10b981"),
        ("负债类", _nz(liabilities), "#ef4444"),
        ("权益类", _nz(equity), "#3b82f6"),
        ("收入类", _nz(revenue), "#9333ea"),
        ("费用类", _nz(expense), "#ea580c"),
    ]

    # 使用网格布局，表格左右排列
    with ui.element("div").classes("tb-grid"):
        for cat_name, cat_items, color in _CATEGORY_CONFIG:
            if not cat_items:
                continue

            cat_open = _sum_field(cat_items, "opening_balance")
            cat_dr = _sum_field(cat_items, "period_debit")
            cat_cr = _sum_field(cat_items, "period_credit")
            cat_close = _sum_field(cat_items, "closing_balance")

            # 每个分类一个卡片
            with ui.card().classes("tb-grid-card"):
                # 卡片头部
                with ui.row().classes("tb-grid-card__header"):
                    with ui.row().classes("items-center gap-2"):
                        ui.element("div").style(f"width:8px;height:8px;border-radius:50%;background:{color}")
                        ui.label(cat_name).classes("tb-grid-card__title")
                        ui.label(f"{len(cat_items)}个科目").classes("tb-grid-card__count")
                    ui.space()
                    ui.label(f"余额: {format_amount(cat_close)}").classes("tb-grid-card__balance")

                # 表格内容
                with ui.column().classes("tb-grid-card__content"):
                    table_html = '<table class="tb-compact-table">'
                    table_html += '<thead><tr>'
                    table_html += '<th class="tb-compact-th">编码</th>'
                    table_html += '<th class="tb-compact-th">名称</th>'
                    table_html += '<th class="tb-compact-th tb-compact-th-num">期初</th>'
                    table_html += '<th class="tb-compact-th tb-compact-th-num">借方</th>'
                    table_html += '<th class="tb-compact-th tb-compact-th-num">贷方</th>'
                    table_html += '<th class="tb-compact-th tb-compact-th-num">期末</th>'
                    table_html += '</tr></thead>'
                    table_html += '<tbody>'

                    for b in cat_items:
                        code = b.get("account_code", "")
                        name = b.get("account_name", "")
                        opening = float(b.get("opening_balance", 0) if b.get("opening_balance") is not None else 0)
                        debit = float(b.get("period_debit", 0) if b.get("period_debit") is not None else 0)
                        credit = float(b.get("period_credit", 0) if b.get("period_credit") is not None else 0)
                        closing = float(b.get("closing_balance", 0) if b.get("closing_balance") is not None else 0)

                        table_html += f'<tr class="tb-compact-row">'
                        table_html += f'<td class="tb-compact-td tb-compact-td-code">{code}</td>'
                        table_html += f'<td class="tb-compact-td tb-compact-td-name">{name}</td>'
                        table_html += f'<td class="tb-compact-td tb-compact-td-num">{format_amount(opening) if opening else "—"}</td>'
                        table_html += f'<td class="tb-compact-td tb-compact-td-num tb-td-debit">{format_amount(debit) if debit else "—"}</td>'
                        table_html += f'<td class="tb-compact-td tb-compact-td-num tb-td-credit">{format_amount(credit) if credit else "—"}</td>'
                        table_html += f'<td class="tb-compact-td tb-compact-td-num" style="font-weight:600">{format_amount(closing) if closing else "—"}</td>'
                        table_html += '</tr>'

                    # 小计行
                    table_html += f'<tr class="tb-compact-row tb-compact-row-subtotal">'
                    table_html += f'<td class="tb-compact-td" colspan="2"><strong>小计</strong></td>'
                    table_html += f'<td class="tb-compact-td tb-compact-td-num">{format_amount(cat_open)}</td>'
                    table_html += f'<td class="tb-compact-td tb-compact-td-num tb-td-debit">{format_amount(cat_dr)}</td>'
                    table_html += f'<td class="tb-compact-td tb-compact-td-num tb-td-credit">{format_amount(cat_cr)}</td>'
                    table_html += f'<td class="tb-compact-td tb-compact-td-num" style="font-weight:700">{format_amount(cat_close)}</td>'
                    table_html += '</tr>'

                    table_html += '</tbody></table>'

                    with ui.card_section().classes("p-0"):
                        ui.html(table_html, sanitize=False)

    # ── 合计行 ──
    with ui.row().classes("tb-total-row"):
        ui.label("全部科目合计").classes("font-bold text-sm")
        with ui.row().classes("gap-6"):
            ui.label(f"期初 {format_amount(total_opening)}")
            ui.label(f"借方 {format_amount(total_debit)}")
            ui.label(f"贷方 {format_amount(total_credit)}")
            ui.label(f"期末 {format_amount(total_closing)}")


def _export_excel(ledger_id, year, month):
    """导出科目余额表为 Excel"""
    try:
        import os
        from datetime import datetime
        balances = ReportService.get_account_balances(ledger_id, year, month)
        if not balances:
            show_toast("无数据可导出", "warning")
            return

        export_dir = os.path.join(os.path.dirname(__file__), "..", "exports")
        os.makedirs(export_dir, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fpath = os.path.join(export_dir, f"科目余额表_{year}{month:02d}_{ts}.xlsx")

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "科目余额表"

        headers = ["科目代码", "科目名称", "类别", "期初余额", "本期借方", "本期贷方", "期末余额"]
        ws.append(headers)

        thin = Side(style="thin", color="E5E7EB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        for b in balances:
            ws.append([
                b.get("account_code", ""),
                b.get("account_name", ""),
                b.get("category", ""),
                float(b.get("opening_balance", 0) if b.get("opening_balance") is not None else 0),
                float(b.get("period_debit", 0) if b.get("period_debit") is not None else 0),
                float(b.get("period_credit", 0) if b.get("period_credit") is not None else 0),
                float(b.get("closing_balance", 0) if b.get("closing_balance") is not None else 0),
            ])

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border

        wb.save(fpath)
        show_toast(f"导出成功: {fpath}", "success")
    except Exception as e:
        show_toast(f"导出失败: {e}", "error")
